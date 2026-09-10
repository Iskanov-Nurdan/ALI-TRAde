"""Выгрузка отчётов в Excel (.xlsx) и CSV."""
import csv
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Символы, с которых Excel и LibreOffice начинают трактовать ячейку как формулу
FORMULA_PREFIXES = ("=", "+", "-", "@", chr(9), chr(13))


def sanitize(value):
    """Обезвреживает значение, которое иначе стало бы формулой в таблице."""
    if not isinstance(value, str):
        return value
    if value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


HEADER_FILL = PatternFill("solid", fgColor="141821")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CELL_BORDER = Border(
    left=Side(style="thin", color="D9DDE3"),
    right=Side(style="thin", color="D9DDE3"),
    top=Side(style="thin", color="D9DDE3"),
    bottom=Side(style="thin", color="D9DDE3"),
)


def _filename(prefix: str, extension: str) -> str:
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    return f"{prefix}-{stamp}.{extension}"


def csv_response(prefix: str, header: list[str], rows) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(prefix, "csv")}"'
    response.write("﻿")  # BOM, чтобы Excel корректно открыл кириллицу
    writer = csv.writer(response, delimiter=";")
    writer.writerow(header)
    writer.writerows([sanitize(value) for value in row] for row in rows)
    return response


def xlsx_response(prefix: str, header: list[str], rows, sheet_title: str = "Отчёт") -> HttpResponse:
    """Формирует книгу Excel: закреплённая шапка, автофильтр, подобранная ширина колонок."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]

    sheet.append(header)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER

    widths = [len(str(title)) + 2 for title in header]
    for row in rows:
        values = [sanitize(value) for value in row]
        sheet.append(values)
        for index, value in enumerate(values):
            if index < len(widths):
                widths[index] = max(widths[index], min(len(str(value if value is not None else "")) + 2, 48))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{sheet.max_row}"
    sheet.row_dimensions[1].height = 26

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{_filename(prefix, "xlsx")}"'
    return response


def export_response(request, prefix: str, header: list[str], rows, sheet_title: str = "Отчёт") -> HttpResponse:
    """Excel по умолчанию, CSV — по параметру ?ext=csv.

    Имя параметра именно `ext`: `format` занят механизмом согласования форматов DRF.
    """
    if (request.query_params.get("ext") or "xlsx").lower() == "csv":
        return csv_response(prefix, header, rows)
    return xlsx_response(prefix, header, rows, sheet_title)
